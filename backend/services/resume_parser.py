"""
Resume Parser Service

This module handles parsing of PDF and DOCX resume files and extracting
structured information using AI.
"""

import os
import json
import logging
from typing import Dict, Any, Optional
import pdfplumber
from docx import Document
from langchain_core.messages import HumanMessage, SystemMessage
from langchain.chat_models import init_chat_model

# Configure logging
logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file using pdfplumber.

    Args:
        file_path: Path to the PDF file

    Returns:
        Extracted text as string

    Raises:
        Exception: If PDF extraction fails
    """
    logger.info(f"📄 Extracting text from PDF: {file_path}")
    try:
        text = ""
        with pdfplumber.open(file_path) as pdf:
            logger.info(f"   PDF has {len(pdf.pages)} pages")
            for i, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
                    logger.debug(f"   Page {i}: {len(page_text)} characters extracted")

        extracted_length = len(text.strip())
        logger.info(f"✅ PDF text extraction complete: {extracted_length} characters total")
        return text.strip()
    except Exception as e:
        logger.error(f"❌ Failed to extract text from PDF: {str(e)}")
        raise Exception(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file using python-docx.

    Args:
        file_path: Path to the DOCX file

    Returns:
        Extracted text as string

    Raises:
        Exception: If DOCX extraction fails
    """
    logger.info(f"📄 Extracting text from DOCX: {file_path}")
    try:
        doc = Document(file_path)
        text = ""

        logger.info(f"   DOCX has {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables")

        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"

        extracted_length = len(text.strip())
        logger.info(f"✅ DOCX text extraction complete: {extracted_length} characters total")
        return text.strip()
    except Exception as e:
        logger.error(f"❌ Failed to extract text from DOCX: {str(e)}")
        raise Exception(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_resume(file_path: str) -> str:
    """
    Extract text from resume file (PDF or DOCX).

    Args:
        file_path: Path to the resume file

    Returns:
        Extracted text as string

    Raises:
        Exception: If file format is unsupported or extraction fails
    """
    if not os.path.exists(file_path):
        raise Exception(f"File not found: {file_path}")

    file_extension = file_path.lower().split('.')[-1]

    if file_extension == 'pdf':
        return extract_text_from_pdf(file_path)
    elif file_extension == 'docx':
        return extract_text_from_docx(file_path)
    else:
        raise Exception(f"Unsupported file format: {file_extension}")


def parse_resume_with_ai(resume_text: str, api_key: Optional[str] = None) -> Dict[str, Any]:
    """
    Use LangChain with AI models to parse resume text and extract structured data with confidence scores.

    Args:
        resume_text: Extracted text from resume
        api_key: API key (if None, will try to get from environment based on model provider)

    Returns:
        Dictionary containing extracted data and confidence scores

    Raises:
        Exception: If AI parsing fails
    """
    logger.info("🤖 Starting AI-based resume parsing with LangChain")

    # Get API key and determine model provider from environment
    openai_key = api_key or os.getenv('OPENAI_API_KEY')
    anthropic_key = os.getenv('ANTHROPIC_API_KEY')

    # Determine which model to use based on available API keys
    if openai_key:
        model_provider = "openai"
        model_name = "gpt-4-turbo-preview"
        os.environ['OPENAI_API_KEY'] = openai_key
        logger.info(f"   Using OpenAI model: {model_name}")
    elif anthropic_key:
        model_provider = "anthropic"
        model_name = "claude-3-5-sonnet-20241022"
        logger.info(f"   Using Anthropic model: {model_name}")
    else:
        logger.error("❌ No API key found in environment variables")
        raise Exception("No API key found. Please set OPENAI_API_KEY or ANTHROPIC_API_KEY in environment variables.")

    logger.info(f"   Resume text length: {len(resume_text)} characters")

    try:
        # Initialize chat model using LangChain's init_chat_model
        logger.info(f"   Initializing {model_provider} chat model...")
        model = init_chat_model(
            model=model_name,
            model_provider=model_provider,
            temperature=0.1,  # Low temperature for more consistent results
        )

        # Add structured output for JSON formatting (OpenAI only)
        if model_provider == "openai":
            model = model.bind(response_format={"type": "json_object"})

        # Create a detailed prompt for structured data extraction
        system_message = SystemMessage(
            content="You are an expert resume parser that extracts structured data from resumes with high accuracy."
        )

        user_prompt = f"""You are an expert resume parser. Extract the following information from the resume text below and return it as a JSON object. For each field, provide a confidence score (0.0 to 1.0) indicating how certain you are about the extracted value.

Resume Text:
{resume_text}

Extract the following fields:
1. name: Full name of the candidate
2. email: Email address
3. phone: Phone number (preferably in Indian format if applicable)
4. company: Current or most recent company name
5. designation: Current or most recent job title/designation
6. skills: Array of technical skills (programming languages, frameworks, tools)
7. experience_years: Total years of professional experience (as integer, best estimate)
8. education: Array of educational qualifications (degrees, institutions)

Return the data in the following JSON format:
{{
    "name": "extracted name or null",
    "email": "extracted email or null",
    "phone": "extracted phone or null",
    "company": "extracted company or null",
    "designation": "extracted designation or null",
    "skills": ["skill1", "skill2", ...] or [],
    "experience_years": integer or null,
    "education": ["degree1", "degree2", ...] or [],
    "confidence_scores": {{
        "name": 0.0-1.0,
        "email": 0.0-1.0,
        "phone": 0.0-1.0,
        "company": 0.0-1.0,
        "designation": 0.0-1.0,
        "skills": 0.0-1.0,
        "experience_years": 0.0-1.0,
        "education": 0.0-1.0
    }}
}}

Important:
- If a field cannot be found, set it to null (or empty array for array fields)
- Confidence scores should reflect how certain you are about the extraction
- For skills, extract only technical skills relevant to the candidate's profession
- For experience_years, calculate based on work history dates if available
- Return ONLY the JSON object, no additional text"""

        user_message = HumanMessage(content=user_prompt)

        # Invoke the model
        logger.info(f"   Calling {model_provider} API via LangChain...")
        response = model.invoke([system_message, user_message])

        logger.info("   ✅ Received response from AI model")

        # Parse the response
        result_text = response.content
        logger.info("   Parsing JSON response...")
        parsed_data = json.loads(result_text)

        # Validate and clean the data
        cleaned_data = {
            'name': parsed_data.get('name'),
            'email': parsed_data.get('email'),
            'phone': parsed_data.get('phone'),
            'company': parsed_data.get('company'),
            'designation': parsed_data.get('designation'),
            'skills': parsed_data.get('skills', []),
            'experience_years': parsed_data.get('experience_years'),
            'education': parsed_data.get('education', []),
            'confidence_scores': parsed_data.get('confidence_scores', {})
        }

        # Ensure confidence scores exist for all fields
        default_scores = {
            'name': 0.5,
            'email': 0.5,
            'phone': 0.5,
            'company': 0.5,
            'designation': 0.5,
            'skills': 0.5,
            'experience_years': 0.5,
            'education': 0.5
        }

        if not cleaned_data['confidence_scores']:
            cleaned_data['confidence_scores'] = default_scores
        else:
            # Fill in missing confidence scores
            for field in default_scores:
                if field not in cleaned_data['confidence_scores']:
                    cleaned_data['confidence_scores'][field] = default_scores[field]

        # Log extracted data summary
        logger.info("   📊 Extraction Summary:")
        logger.info(f"      Name: {cleaned_data.get('name') or 'Not found'}")
        logger.info(f"      Email: {cleaned_data.get('email') or 'Not found'}")
        logger.info(f"      Phone: {cleaned_data.get('phone') or 'Not found'}")
        logger.info(f"      Company: {cleaned_data.get('company') or 'Not found'}")
        logger.info(f"      Designation: {cleaned_data.get('designation') or 'Not found'}")
        logger.info(f"      Skills: {len(cleaned_data.get('skills', []))} found")
        logger.info(f"      Experience: {cleaned_data.get('experience_years') or 'Not found'} years")
        logger.info(f"      Education: {len(cleaned_data.get('education', []))} entries found")

        avg_confidence = sum(cleaned_data['confidence_scores'].values()) / len(cleaned_data['confidence_scores'])
        logger.info(f"   📈 Average confidence score: {avg_confidence:.2%}")

        logger.info("✅ AI parsing completed successfully")

        return cleaned_data

    except json.JSONDecodeError as e:
        logger.error(f"Failed to parse AI response as JSON: {str(e)}")
        raise Exception(f"Failed to parse AI response as JSON: {str(e)}")
    except Exception as e:
        logger.error(f"AI parsing failed: {str(e)}")
        raise Exception(f"AI parsing failed: {str(e)}")


def parse_resume(file_path: str, api_key: Optional[str] = None) -> Dict[str, Any]:
    """
    Parse a resume file and extract structured data using AI via LangChain.

    This is the main entry point for resume parsing. It:
    1. Extracts text from PDF/DOCX files
    2. Uses AI (OpenAI/Anthropic via LangChain) to parse and extract structured data
    3. Calculates confidence scores for each field
    4. Returns structured data

    Args:
        file_path: Path to the resume file (PDF or DOCX)
        api_key: Optional API key (if None, uses environment variable OPENAI_API_KEY or ANTHROPIC_API_KEY)

    Returns:
        Dictionary containing extracted data and confidence scores with keys:
        - name: str or None
        - email: str or None
        - phone: str or None
        - company: str or None
        - designation: str or None
        - skills: list of str
        - experience_years: int or None
        - education: list of str
        - confidence_scores: dict with scores for each field

    Raises:
        Exception: If parsing fails at any stage
    """
    try:
        # Step 1: Extract text from the file
        logger.info("Step 1/2: Text Extraction")
        resume_text = extract_text_from_resume(file_path)

        print(resume_text)

        if not resume_text or len(resume_text.strip()) < 50:
            logger.error(f"Insufficient text extracted ({len(resume_text.strip())} characters)")
            raise Exception("Insufficient text extracted from resume. File may be empty or corrupted.")

        # TODO: Remove PII from resume_text
        
        logger.info("")
        logger.info("Step 2/2: AI Parsing")

        # Step 2: Parse with AI
        extracted_data = parse_resume_with_ai(resume_text, api_key)

        logger.info("")
        logger.info("="*60)
        logger.info("Resume parsing completed successfully!")
        logger.info("="*60)

        return extracted_data

    except Exception as e:
        logger.error("")
        logger.error("="*60)
        logger.error(f"Resume parsing failed: {str(e)}")
        logger.error("="*60)
        # Re-raise with context
        raise Exception(f"Resume parsing failed: {str(e)}")